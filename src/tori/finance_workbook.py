"""Human-readable Excel implementation of the FinanceRepository boundary."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
import hashlib
import os
from pathlib import Path
import secrets
from typing import Callable, Iterable, Sequence, TypeVar

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .finance import (
    CONTRACT_VERSION, Bill, Budget, Debt, FinanceConflictError, FinanceError,
    FinanceSnapshot, FinanceValidationError, Goal, MerchantRule, Transaction, money,
)


WORKBOOK_NAME = "tori_finance.xlsx"
SHEET_NAMES = (
    "Transactions", "Merchant Rules", "Bills", "Debts", "Budget", "Goals", "Summary"
)
_MAX_ROWS = 100_000

TRANSACTION_HEADERS = (
    "Transaction ID", "Date", "Merchant", "Original Description", "Amount",
    "Transaction Type", "Category", "Account Label", "Import Source", "Import Batch ID",
    "Source Transaction ID", "Duplicate Fingerprint", "Related Bill ID", "Related Debt ID", "Notes",
)
RULE_HEADERS = (
    "Rule ID", "Match Type", "Match Text", "Account Scope", "Normalized Merchant",
    "Category", "Priority", "Active", "Notes",
)
BILL_HEADERS = (
    "Bill ID", "Name", "Category", "Amount Kind", "Expected Amount", "Rolling Average",
    "Next Due Date", "Frequency", "Due Day", "Autopay", "Active", "Notes",
)
DEBT_HEADERS = (
    "Debt ID", "Name", "Debt Type", "Current Balance", "APR", "Minimum Payment",
    "Target Payment", "Next Due Date", "Due Day", "Priority", "Balance As Of", "Active", "Notes",
)
BUDGET_HEADERS = (
    "Month", "Expected Income", "Starting Available Funds", "Essential Spending Target",
    "Discretionary Target", "Savings Target", "Extra Debt Payment Target", "Reserve Buffer", "Notes",
)
GOAL_HEADERS = (
    "Goal ID", "Name", "Goal Type", "Target Amount", "Current Amount", "Linked Debt ID",
    "Debt Starting Balance", "Target Date", "Priority", "Status", "Amount As Of", "Notes",
)
SUMMARY_HEADERS = ("Metric", "Value")


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as exc:
        raise FinanceError("The Finance workbook could not be read.", code="finance_unavailable") from exc
    return f"sha256:{digest.hexdigest()}"


def _optional_text(value: object) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    return str(value)


def _required_text(value: object, label: str) -> str:
    result = _optional_text(value)
    if result is None:
        raise FinanceValidationError(f"{label} is required in the workbook.")
    return result


def _date_value(value: object, label: str, *, optional: bool = False) -> date | None:
    if value is None or value == "":
        if optional:
            return None
        raise FinanceValidationError(f"{label} is required in the workbook.")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise FinanceValidationError(f"{label} must be YYYY-MM-DD.") from exc


def _bool_value(value: object, label: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().casefold() in {"true", "false"}:
        return value.strip().casefold() == "true"
    raise FinanceValidationError(f"{label} must be TRUE or FALSE.")


def _int_value(value: object, label: str, *, optional: bool = False) -> int | None:
    if value is None or value == "":
        if optional:
            return None
        raise FinanceValidationError(f"{label} is required in the workbook.")
    if isinstance(value, bool):
        raise FinanceValidationError(f"{label} must be an integer.")
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise FinanceValidationError(f"{label} must be an integer.") from exc
    if isinstance(value, float) and value != result:
        raise FinanceValidationError(f"{label} must be an integer.")
    return result


def _money_value(value: object, label: str, *, optional: bool = False) -> Decimal | None:
    if value is None or value == "":
        if optional:
            return None
        raise FinanceValidationError(f"{label} is required in the workbook.")
    return money(Decimal(str(value)))


def _rows(sheet: Worksheet, headers: tuple[str, ...]) -> Iterable[tuple[object, ...]]:
    if sheet.max_row > _MAX_ROWS + 1:
        raise FinanceValidationError(f"{sheet.title} exceeds the Finance V1 row limit.")
    actual = tuple(cell.value for cell in sheet[1])
    if actual != headers:
        raise FinanceValidationError(f"{sheet.title} workbook headers do not match Finance V1.")
    for values in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True):
        if any(value is not None and value != "" for value in values):
            yield values


class WorkbookFinanceRepository:
    """Finance workbook adapter with optimistic revisions and atomic replacement."""

    def __init__(self, data_root: Path | str, *, prohibited_roots: Sequence[Path | str] = ()) -> None:
        root = Path(data_root).expanduser().absolute()
        self._validate_path(root, tuple(Path(item).expanduser().absolute() for item in prohibited_roots))
        self._root = root
        self._workbook_path = root / WORKBOOK_NAME

    @property
    def workbook_path(self) -> Path:
        return self._workbook_path

    @property
    def data_root(self) -> Path:
        return self._root

    @staticmethod
    def _validate_path(root: Path, prohibited: tuple[Path, ...]) -> None:
        if not root.is_absolute():
            raise FinanceValidationError("Finance data root must be an absolute path.")
        for item in (root, *root.parents):
            if item.exists() and item.is_symlink():
                raise FinanceValidationError("Finance data root cannot resolve through a symbolic link.")
        resolved = root.resolve(strict=False)
        for blocked in prohibited:
            blocked_resolved = blocked.resolve(strict=False)
            if resolved == blocked_resolved or blocked_resolved in resolved.parents:
                raise FinanceValidationError("Finance data root must remain outside protected Tori storage.")

    def exists(self) -> bool:
        if self._root.is_symlink() or self._workbook_path.is_symlink():
            raise FinanceError("Finance storage changed to an unsafe symbolic-link path.", code="finance_unsafe_path")
        return self._workbook_path.is_file()

    def revision(self) -> str | None:
        return _digest(self._workbook_path) if self.exists() else None

    def create(self, *, currency: str, expected_revision: str | None = None) -> FinanceSnapshot:
        if self.revision() != expected_revision:
            raise FinanceConflictError()
        if expected_revision is not None or self.exists():
            raise FinanceConflictError("Finance workbook already exists; nothing was overwritten.")
        try:
            self._root.mkdir(mode=0o700, parents=True, exist_ok=True)
            (self._root / "imports" / "incoming").mkdir(mode=0o700, parents=True, exist_ok=True)
            (self._root / "imports" / "processed").mkdir(mode=0o700, parents=True, exist_ok=True)
        except OSError as exc:
            raise FinanceError("The Finance data root could not be created.", code="finance_unavailable") from exc
        snapshot = FinanceSnapshot("", currency=currency)
        self._atomic_write(snapshot)
        return self.read()

    def read(self) -> FinanceSnapshot:
        if not self.exists():
            raise FinanceError("Finance has not been initialized.", code="finance_not_initialized")
        revision = self.revision()
        assert revision is not None
        try:
            workbook = load_workbook(self._workbook_path, data_only=True, read_only=True)
        except Exception as exc:
            raise FinanceError("The Finance workbook is unavailable or invalid.", code="finance_invalid_workbook") from exc
        try:
            if tuple(workbook.sheetnames) != SHEET_NAMES:
                raise FinanceValidationError("Finance workbook sheets do not match the V1 contract.")
            summary = workbook["Summary"]
            metadata = {str(row[0]): row[1] for row in _rows(summary, SUMMARY_HEADERS)}
            if str(metadata.get("Finance Contract Version", "")) != CONTRACT_VERSION:
                raise FinanceValidationError("Finance workbook contract version is unsupported.")
            currency = _required_text(metadata.get("Currency"), "Summary Currency")
            transactions = tuple(self._transaction(row) for row in _rows(workbook["Transactions"], TRANSACTION_HEADERS))
            rules = tuple(self._rule(row) for row in _rows(workbook["Merchant Rules"], RULE_HEADERS))
            bills = tuple(self._bill(row) for row in _rows(workbook["Bills"], BILL_HEADERS))
            debts = tuple(self._debt(row) for row in _rows(workbook["Debts"], DEBT_HEADERS))
            budgets = tuple(self._budget(row) for row in _rows(workbook["Budget"], BUDGET_HEADERS))
            goals = tuple(self._goal(row) for row in _rows(workbook["Goals"], GOAL_HEADERS))
            projection = tuple(
                (str(key), "" if value is None else str(value))
                for key, value in metadata.items()
                if key not in {"Finance Contract Version", "Currency"}
            )
            return FinanceSnapshot(revision, currency, transactions, rules, bills, debts, budgets, goals, projection)
        finally:
            workbook.close()

    def replace(self, snapshot: FinanceSnapshot, *, expected_revision: str) -> FinanceSnapshot:
        if self.revision() != expected_revision:
            raise FinanceConflictError()
        validated = FinanceSnapshot(
            expected_revision, snapshot.currency, snapshot.transactions, snapshot.merchant_rules,
            snapshot.bills, snapshot.debts, snapshot.budgets, snapshot.goals, snapshot.summary_rows,
        )
        self._atomic_write(validated)
        result = self.read()
        expected_without_revision = validated.with_revision(result.revision)
        if result != expected_without_revision:
            raise FinanceError("Finance workbook verification failed after writing.", code="finance_verification_failed")
        return result

    def _atomic_write(self, snapshot: FinanceSnapshot) -> None:
        temp = self._root / f".{WORKBOOK_NAME}.{secrets.token_hex(8)}.tmp"
        workbook = self._build_workbook(snapshot)
        try:
            workbook.save(temp)
            with temp.open("rb") as written:
                os.fsync(written.fileno())
            os.replace(temp, self._workbook_path)
            try:
                directory_fd = os.open(self._root, os.O_RDONLY)
                try:
                    os.fsync(directory_fd)
                finally:
                    os.close(directory_fd)
            except OSError:
                pass
        except OSError as exc:
            raise FinanceError("Finance workbook could not be written atomically.", code="finance_unavailable") from exc
        finally:
            workbook.close()
            try:
                temp.unlink(missing_ok=True)
            except OSError:
                pass

    @staticmethod
    def _build_workbook(snapshot: FinanceSnapshot) -> Workbook:
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = {name: workbook.create_sheet(name) for name in SHEET_NAMES}
        WorkbookFinanceRepository._write_sheet(sheets["Transactions"], TRANSACTION_HEADERS, (
            (item.identifier, item.date, item.merchant, item.original_description, item.amount, item.transaction_type,
             item.category, item.account_label, item.import_source, item.import_batch_id, item.source_transaction_id,
             item.duplicate_fingerprint, item.related_bill_id, item.related_debt_id, item.notes)
            for item in snapshot.transactions
        ), money_columns={5})
        WorkbookFinanceRepository._write_sheet(sheets["Merchant Rules"], RULE_HEADERS, (
            (item.identifier, item.match_type, item.match_text, item.account_scope, item.normalized_merchant,
             item.category, item.priority, item.active, item.notes) for item in snapshot.merchant_rules
        ))
        WorkbookFinanceRepository._write_sheet(sheets["Bills"], BILL_HEADERS, (
            (item.identifier, item.name, item.category, item.amount_kind, item.expected_amount, item.rolling_average,
             item.next_due_date, item.frequency, item.due_day, item.autopay, item.active, item.notes) for item in snapshot.bills
        ), money_columns={5, 6})
        WorkbookFinanceRepository._write_sheet(sheets["Debts"], DEBT_HEADERS, (
            (item.identifier, item.name, item.debt_type, item.current_balance, item.apr, item.minimum_payment,
             item.target_payment, item.next_due_date, item.due_day, item.priority, item.balance_as_of, item.active, item.notes)
            for item in snapshot.debts
        ), money_columns={4, 6, 7})
        WorkbookFinanceRepository._write_sheet(sheets["Budget"], BUDGET_HEADERS, (
            (item.month, item.expected_income, item.starting_available_funds, item.essential_spending_target,
             item.discretionary_target, item.savings_target, item.extra_debt_payment_target,
             item.reserve_buffer, item.notes) for item in snapshot.budgets
        ), money_columns={2, 3, 4, 5, 6, 7, 8})
        WorkbookFinanceRepository._write_sheet(sheets["Goals"], GOAL_HEADERS, (
            (item.identifier, item.name, item.goal_type, item.target_amount, item.current_amount,
             item.linked_debt_id, item.debt_starting_balance, item.target_date, item.priority,
             item.status, item.amount_as_of, item.notes) for item in snapshot.goals
        ), money_columns={4, 5, 7})
        summary_rows = (
            ("Finance Contract Version", CONTRACT_VERSION), ("Currency", snapshot.currency), *snapshot.summary_rows
        )
        WorkbookFinanceRepository._write_sheet(sheets["Summary"], SUMMARY_HEADERS, summary_rows)
        for sheet in sheets.values():
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        return workbook

    @staticmethod
    def _write_sheet(
        sheet: Worksheet, headers: tuple[str, ...], rows: Iterable[Sequence[object]], *, money_columns: set[int] | None = None
    ) -> None:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="3B4A66")
        for row in rows:
            prepared = tuple(float(value) if isinstance(value, Decimal) else value for value in row)
            sheet.append(prepared)
            for cell, original in zip(sheet[sheet.max_row], row):
                if isinstance(original, str):
                    # User/source text remains text even when it begins with a
                    # spreadsheet formula marker.
                    cell.data_type = "s"
        for index, header in enumerate(headers, start=1):
            width = max(
                [len(header)]
                + [len(str(sheet.cell(row, index).value or "")) for row in range(2, min(sheet.max_row, 200) + 1)]
            )
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(width + 2, 12), 42)
        for column in money_columns or set():
            for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
                for item in cell:
                    item.number_format = '$#,##0.00;[Red]-$#,##0.00'

    @staticmethod
    def _transaction(v: tuple[object, ...]) -> Transaction:
        return Transaction(
            _required_text(v[0], "Transaction ID"), _date_value(v[1], "Transaction Date"),
            _required_text(v[2], "Merchant"), _required_text(v[3], "Original Description"),
            _money_value(v[4], "Transaction Amount"), _required_text(v[5], "Transaction Type"),
            _required_text(v[6], "Category"), _required_text(v[7], "Account Label"),
            _required_text(v[8], "Import Source"), _optional_text(v[9]), _optional_text(v[10]),
            _required_text(v[11], "Duplicate Fingerprint"), _optional_text(v[12]), _optional_text(v[13]),
            _optional_text(v[14]) or "",
        )

    @staticmethod
    def _rule(v: tuple[object, ...]) -> MerchantRule:
        return MerchantRule(
            _required_text(v[0], "Rule ID"), _required_text(v[1], "Match Type"),
            _required_text(v[2], "Match Text"), _required_text(v[4], "Normalized Merchant"),
            _required_text(v[5], "Category"), _int_value(v[6], "Priority"), _bool_value(v[7], "Active"),
            _optional_text(v[3]), _optional_text(v[8]) or "",
        )

    @staticmethod
    def _bill(v: tuple[object, ...]) -> Bill:
        return Bill(
            _required_text(v[0], "Bill ID"), _required_text(v[1], "Bill name"), _required_text(v[2], "Bill category"),
            _required_text(v[3], "Amount Kind"), _money_value(v[4], "Expected Amount"), _date_value(v[6], "Next Due Date"),
            _required_text(v[7], "Frequency"), _bool_value(v[9], "Autopay"), _bool_value(v[10], "Active"),
            _money_value(v[5], "Rolling Average", optional=True), _int_value(v[8], "Due Day", optional=True), _optional_text(v[11]) or "",
        )

    @staticmethod
    def _debt(v: tuple[object, ...]) -> Debt:
        return Debt(
            _required_text(v[0], "Debt ID"), _required_text(v[1], "Debt name"), _required_text(v[2], "Debt Type"),
            _money_value(v[3], "Current Balance"), Decimal(str(v[4])), _money_value(v[5], "Minimum Payment"),
            _money_value(v[6], "Target Payment"), _date_value(v[7], "Next Due Date"), _int_value(v[9], "Priority"),
            _date_value(v[10], "Balance As Of"), _bool_value(v[11], "Active"), _int_value(v[8], "Due Day", optional=True),
            _optional_text(v[12]) or "",
        )

    @staticmethod
    def _budget(v: tuple[object, ...]) -> Budget:
        return Budget(
            _required_text(v[0], "Budget Month"), _money_value(v[1], "Expected Income"),
            _money_value(v[3], "Essential Spending Target"), _money_value(v[4], "Discretionary Target"),
            _money_value(v[5], "Savings Target"), _money_value(v[6], "Extra Debt Payment Target"),
            _money_value(v[7], "Reserve Buffer"), _money_value(v[2], "Starting Available Funds", optional=True),
            _optional_text(v[8]) or "",
        )

    @staticmethod
    def _goal(v: tuple[object, ...]) -> Goal:
        return Goal(
            _required_text(v[0], "Goal ID"), _required_text(v[1], "Goal name"), _required_text(v[2], "Goal Type"),
            _money_value(v[3], "Target Amount"), _int_value(v[8], "Priority"), _required_text(v[9], "Goal Status"),
            _date_value(v[10], "Amount As Of"), _money_value(v[4], "Current Amount", optional=True),
            _optional_text(v[5]), _money_value(v[6], "Debt Starting Balance", optional=True),
            _date_value(v[7], "Target Date", optional=True), _optional_text(v[11]) or "",
        )
