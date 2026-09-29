from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from tori.finance import FinanceConflictError, FinanceValidationError
from tori.finance_workbook import SHEET_NAMES, WorkbookFinanceRepository

from finance_fixtures import synthetic_snapshot


class WorkbookFinanceRepositoryTests(unittest.TestCase):
    def test_creates_human_readable_seven_sheet_workbook_and_round_trips(self) -> None:
        with TemporaryDirectory() as temporary:
            repository = WorkbookFinanceRepository(Path(temporary) / "finance")
            created = repository.create(currency="USD")
            saved = repository.replace(synthetic_snapshot(created.revision), expected_revision=created.revision)
            self.assertEqual(repository.read(), saved)
            workbook = load_workbook(repository.workbook_path, read_only=True)
            try:
                self.assertEqual(tuple(workbook.sheetnames), SHEET_NAMES)
                self.assertEqual(workbook["Transactions"]["A1"].value, "Transaction ID")
                self.assertEqual(workbook["Debts"]["D2"].value, 3000)
                self.assertEqual(workbook["Summary"]["A2"].value, "Finance Contract Version")
            finally:
                workbook.close()

    def test_manual_edit_changes_revision_and_rejects_stale_write(self) -> None:
        with TemporaryDirectory() as temporary:
            repository = WorkbookFinanceRepository(Path(temporary) / "finance")
            created = repository.create(currency="USD")
            stale = created.revision
            workbook = load_workbook(repository.workbook_path)
            workbook["Summary"]["B4"] = "manual note"
            workbook.save(repository.workbook_path)
            workbook.close()
            self.assertNotEqual(repository.revision(), stale)
            with self.assertRaises(FinanceConflictError):
                repository.replace(created, expected_revision=stale)

    def test_atomic_write_leaves_no_temporary_files(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary) / "finance"
            repository = WorkbookFinanceRepository(root)
            created = repository.create(currency="USD")
            repository.replace(replace(created, summary_rows=(("Note", "synthetic"),)), expected_revision=created.revision)
            self.assertEqual([item.name for item in root.iterdir() if item.name.startswith(".tori_finance")], [])

    def test_statement_text_that_looks_like_formula_remains_literal_text(self) -> None:
        with TemporaryDirectory() as temporary:
            repository = WorkbookFinanceRepository(Path(temporary) / "finance")
            created = repository.create(currency="USD")
            snapshot = synthetic_snapshot(created.revision)
            transaction = replace(snapshot.transactions[0], original_description="=HYPERLINK(\"bad\")")
            repository.replace(replace(snapshot, transactions=(transaction,) + snapshot.transactions[1:]), expected_revision=created.revision)
            workbook = load_workbook(repository.workbook_path, data_only=False, read_only=True)
            try:
                cell = workbook["Transactions"]["D2"]
                self.assertEqual(cell.value, '=HYPERLINK("bad")')
                self.assertEqual(cell.data_type, "s")
            finally:
                workbook.close()

    def test_rejects_protected_and_symlinked_roots(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            with self.assertRaises(FinanceValidationError):
                WorkbookFinanceRepository(root / "repo" / "finance", prohibited_roots=(root / "repo",))
            actual = root / "actual"
            actual.mkdir()
            link = root / "link"
            link.symlink_to(actual, target_is_directory=True)
            with self.assertRaises(FinanceValidationError):
                WorkbookFinanceRepository(link / "finance")


if __name__ == "__main__":
    unittest.main()
